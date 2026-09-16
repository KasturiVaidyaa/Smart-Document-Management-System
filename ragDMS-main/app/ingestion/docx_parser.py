"""DOCX parsing via python-docx — preserves paragraph/table interleaving."""
from __future__ import annotations
import logging
from pathlib import Path
from docx import Document

logger = logging.getLogger("ragdms.docx_parser")


def _is_heading_style(style_name: str | None) -> bool:
    if not style_name:
        return False
    return style_name.lower().startswith("heading") or style_name.lower() == "title"


def parse_docx(path: Path) -> list[dict]:
    """Extract blocks from a DOCX file preserving document order."""
    doc = Document(str(path))
    blocks: list[dict] = []
    current_heading: str | None = None
    
    # Walk XML body children to preserve paragraph/table interleaving
    body = doc.element.body
    para_map = {p._p: p for p in doc.paragraphs}
    table_map = {t._tbl: t for t in doc.tables}
    
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        
        if tag == "p" and child in para_map:
            p = para_map[child]
            text = (p.text or "").strip()
            if not text:
                continue
            is_heading = _is_heading_style(p.style.name if p.style else None)
            kind = "heading" if is_heading else "body"
            if is_heading:
                current_heading = text
            blocks.append({
                "text": text,
                "page": 1,  # DOCX doesn't have reliable page numbers
                "kind": kind,
                "heading": current_heading,
            })
        
        elif tag == "tbl" and child in table_map:
            t = table_map[child]
            rows = [[(c.text or "").strip() for c in row.cells] for row in t.rows]
            if not rows:
                continue
            for row in rows[1:] if len(rows) > 1 else rows:
                text = " | ".join(row)
                if text.strip():
                    blocks.append({
                        "text": text,
                        "page": 1,
                        "kind": "table_row",
                        "heading": current_heading,
                    })
    
    logger.info("Parsed DOCX: %d blocks", len(blocks))
    return blocks
